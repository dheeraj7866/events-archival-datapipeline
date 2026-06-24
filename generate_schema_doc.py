from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.2)
    section.right_margin  = Cm(2.2)

# ── Colour palette ────────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1F, 0x49, 0x7D)   # heading / header fill
MID_BLUE   = RGBColor(0xBD, 0xD7, 0xEE)   # section header row fill
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)   # alternating row fill
AMBER      = RGBColor(0xFF, 0xC0, 0x00)   # "Confirm" column header
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
RED        = RGBColor(0xC0, 0x00, 0x00)

DARK_BLUE_HEX  = "1F497D"
AMBER_HEX      = "FFC000"

def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    hex_color = str(rgb).upper()          # RGBColor.__str__ returns 6-char hex
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"),   val.get("val",   "single"))
            el.set(qn("w:sz"),    val.get("sz",    "4"))
            el.set(qn("w:space"), val.get("space", "0"))
            el.set(qn("w:color"), val.get("color", "auto"))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def para_font(para, bold=False, size=None, color=None, italic=False):
    for run in para.runs:
        run.bold   = bold
        run.italic = italic
        if size:  run.font.size  = Pt(size)
        if color: run.font.color.rgb = color

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size  = Pt(16)
        run.font.color.rgb = DARK_BLUE
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after  = Pt(4)
        # bottom border
        pPr  = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bot  = OxmlElement("w:bottom")
        bot.set(qn("w:val"),   "single")
        bot.set(qn("w:sz"),    "6")
        bot.set(qn("w:space"), "1")
        bot.set(qn("w:color"), DARK_BLUE_HEX)
        pBdr.append(bot)
        pPr.append(pBdr)
    elif level == 2:
        run.font.size  = Pt(12)
        run.font.color.rgb = DARK_BLUE
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after  = Pt(2)
    return p

def add_note(doc, text, color=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = color or RGBColor(0x59, 0x59, 0x59)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(6)

def make_table(doc, headers, rows, col_widths, confirm_col=None):
    """
    headers    : list of column header strings
    rows       : list of row tuples
    col_widths : list of Inches() widths
    confirm_col: index of the "Confirm" column (gets amber header, bold content)
    """
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"

    # Header row
    hrow = t.rows[0]
    hrow.height = Cm(0.75)
    for i, (cell, hdr) in enumerate(zip(hrow.cells, headers)):
        set_cell_bg(cell, DARK_BLUE if confirm_col is None or i != confirm_col else AMBER)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = WHITE if (confirm_col is None or i != confirm_col) else BLACK

    # Data rows
    for ri, row_data in enumerate(rows):
        r = t.add_row()
        bg = LIGHT_GREY if ri % 2 == 0 else WHITE
        for ci, (cell, val) in enumerate(zip(r.cells, row_data)):
            set_cell_bg(cell, bg)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if confirm_col is not None and ci == confirm_col and val and val != "—":
                run.bold = True
                run.font.color.rgb = RED

    # Column widths
    for i, width in enumerate(col_widths):
        for row in t.rows:
            row.cells[i].width = width

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

# ══════════════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ══════════════════════════════════════════════════════════════════════════════
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(40)
run = title.add_run("Vendor API Event — Field Schema")
run.bold = True
run.font.size = Pt(24)
run.font.color.rgb = DARK_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sub.add_run("Version 1  ·  For Product Approval")
r2.font.size = Pt(12)
r2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.add_paragraph()
box_para = doc.add_paragraph()
box_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = box_para.add_run(
    "Every vendor API call made by identity-api / los-api / payment-api\n"
    "is logged as one row in the vendor archive.\n"
    "This document lists every field stored, its purpose, and open decisions\n"
    "that require product sign-off before dev go-live."
)
r3.font.size = Pt(10)
r3.italic = True
r3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — Call Identity
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 1 — Call Identity")
add_note(doc, "Auto-generated by the system. No product input needed.")

make_table(doc,
    headers=["Field", "Example", "What it stores"],
    rows=[
        ("request_id",      "a1b2c3d4-5e6f-…",          "Unique ID for this vendor call. Prevents duplicate entries if the same event is delivered more than once."),
        ("correlation_id",  "req-abc-123",               "The user's original request ID. Links all vendor calls that happened within a single API request."),
        ("created_at",      "2026-06-01 14:32:00 IST",   "Timestamp when the vendor call was made (Asia/Kolkata)."),
        ("schema_version",  "1",                         "Internal contract versioning. Not user-facing."),
    ],
    col_widths=[Inches(1.6), Inches(1.8), Inches(4.0)],
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — Source
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 2 — Source")
add_note(doc, "Confirm: are these the correct service names and vendor IDs for Grafana dashboards and alerts?")

make_table(doc,
    headers=["Field", "Example", "What it stores", "Confirm with Product"],
    rows=[
        ("service",        "identity-api",        "Which of our APIs made the call.",                                              "Values: identity-api, los-api, payment-api — are these correct and complete?"),
        ("environment",    "staging",              "staging or prod.",                                                             "—"),
        ("vendor_id",      "karza",                "Which vendor was called.",                                                     "Current values: karza, nsdl, easebuzz, crif-highmark, bank-statement-analyser — is this the complete list?"),
        ("endpoint",       "/v2/pan/validate",     "The specific API path called on the vendor.",                                  "—"),
        ("vendor_ref_id",  "EZB-TXN-99887766",    "Vendor's own transaction ID. Optional — only set when the vendor returns one. Used to reconcile against vendor invoices.", "Confirm this is needed. Currently kept for all vendors."),
    ],
    col_widths=[Inches(1.3), Inches(1.5), Inches(2.8), Inches(2.4)],
    confirm_col=3,
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — Loan Context
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 3 — Loan Context")
add_note(doc, "Confirm each field before go-live.")

make_table(doc,
    headers=["Field", "Example", "What it stores", "Confirm with Product"],
    rows=[
        ("loan_lifecycle_stage",    "KYC",              "The stage of the loan journey that triggered this vendor call.",
         "Full list (locked): LEAD → SELFIE → KYC → PAN_AADHAAR_SEED → LOCATION_BRE → BUREAU_BRE → BANK_BRE → REPEAT_BRE → UNDERWRITING → DISBURSED → REPAID → OVERDUE → CLOSED → WRITTEN_OFF — does this cover all stages?"),
        ("loan_application_number", "LAN-2026-001",     "Human-readable loan reference. Optional — only set when available at the time of the call.",
         "Is this available at every stage, or only after loan creation?"),
        ("user_id",                 "USR-9876",         "The borrower's global customer ID.",
         "Confirm this is the correct identifier. What format does it take — UUID, integer, or string?"),
        ("consent_id",              "CONSENT-XYZ-001",  "The consent record ID for this customer. Required for all PII-touching stages (KYC, PAN_AADHAAR_SEED, BUREAU_BRE, BANK_BRE).",
         "Confirm this is a FK to the consent_records table. Is the naming correct?"),
    ],
    col_widths=[Inches(1.5), Inches(1.3), Inches(2.7), Inches(2.5)],
    confirm_col=3,
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — PII Handling
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 4 — PII Handling")
add_note(doc,
    "No raw PAN, mobile number, or Aadhaar is ever stored in the database. "
    "The table below shows exactly how each PII field is handled. "
    "Confirm this approach is acceptable to compliance and product.")

make_table(doc,
    headers=["Field stored", "Original input", "What is stored", "Reversible?"],
    rows=[
        ("pan_masked",              "ABCPK1234F",       "First 3 + **** + last 3  →  ABC****34F",                                                                           "No"),
        ("mobile_hash",             "+919876543210",    "HMAC-SHA256 with a secret salt  →  a3f9c2b1… (64 chars). A customer can be looked up by hashing their number at query time.",   "No"),
        ("mobile_last4",            "+919876543210",    "Last 4 digits stored plain  →  3210",                                                                              "N/A"),
        ("aadhaar_last4_hash",      "3210 (last 4)",    "HMAC-SHA256 with a separate secret salt  →  d8e1f3ca… (64 chars)",                                                 "No"),
        ("aadhaar_last4_encrypted", "—",                "Reserved column — empty until D2 (KMS encryption decision) is resolved",                                           "N/A"),
    ],
    col_widths=[Inches(1.6), Inches(1.4), Inches(3.6), Inches(1.4)],
)

add_note(doc,
    "Raw request and response payloads are stored encrypted in S3 forever for RBI audit. "
    "The database holds a redacted copy for 30 days only, after which it is wiped. "
    "The S3 copy is never deleted (Object Lock COMPLIANCE + Legal Hold).",
    color=RGBColor(0x17, 0x59, 0x0E),
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — Call Outcome
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 5 — Call Outcome")
add_note(doc, "Standard result fields. Confirm the status values cover your alerting and reporting needs.")

make_table(doc,
    headers=["Field", "Example", "What it stores", "Confirm with Product"],
    rows=[
        ("status",        "SUCCESS",                   "Overall result of the vendor call.",
         "4 values: SUCCESS (2xx), FAILURE (4xx/5xx), TIMEOUT (connection/read timeout), NETWORK_ERROR (DNS, refused). Are these the right categories for your alert rules?"),
        ("http_status",   "200, 422, 0",               "HTTP status code returned by the vendor. 0 for timeouts and network errors.",  "—"),
        ("latency_ms",    "143",                       "How long the vendor call took, in milliseconds.",                               "—"),
        ("error_code",    "ECONNABORTED",              "Error code when the call fails (e.g. ETIMEDOUT, NOT_FOUND). Empty on success.", "—"),
        ("error_message", "timeout of 30000ms…",      "Error message when the call fails. Must not contain PII — the calling service is responsible for sanitising.", "—"),
    ],
    col_widths=[Inches(1.3), Inches(1.3), Inches(3.0), Inches(2.4)],
    confirm_col=3,
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 6 — Retention Policy
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 6 — Data Retention")
add_note(doc, "For compliance awareness.")

make_table(doc,
    headers=["Data", "Retention", "Location"],
    rows=[
        ("All metadata fields (Sections 1–5)",              "90 days, then permanently deleted",    "ClickHouse database"),
        ("Request / response payload — redacted copy",      "30 days, then wiped from database",    "ClickHouse database"),
        ("Request / response payload — raw encrypted copy", "Forever (Object Lock + Legal Hold)",   "S3 — KMS encrypted, legally immutable"),
    ],
    col_widths=[Inches(2.8), Inches(2.2), Inches(3.0)],
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 7 — Open Items
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "Section 7 — Open Items for Product Sign-off")
add_note(doc, "These must be confirmed before dev go-live.")

make_table(doc,
    headers=["#", "Question", "Owner"],
    rows=[
        ("1", "Is the loan_lifecycle_stage enum complete? (14 stages listed in Section 3)",                                    "Product"),
        ("2", "Is loan_application_number available at all loan stages, or only after loan creation?",                        "Product"),
        ("3", "Confirm the user_id format and which identifier it maps to in your system",                                    "Product / Engineering"),
        ("4", "Is the 4-way status split (SUCCESS / FAILURE / TIMEOUT / NETWORK_ERROR) sufficient for alert rules?",          "Product / Engineering"),
        ("5", "Confirm vendor_id values are the complete list required on Grafana dashboards",                                 "Product"),
        ("6", "Is PII handling in Section 4 acceptable? (No raw PAN / mobile / Aadhaar ever stored in the database)",        "Product / Compliance"),
        ("7", "90-day row retention and 30-day payload retention — are these acceptable for operational queries?",            "Product / Compliance"),
    ],
    col_widths=[Inches(0.3), Inches(6.4), Inches(1.3)],
)

# ══════════════════════════════════════════════════════════════════════════════
# Footer note
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
footer_para = doc.add_paragraph()
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer_para.add_run("Finagle / Tez Credit  ·  Vendor Archive  ·  Schema v1  ·  2026-06-01")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

# ── Save ──────────────────────────────────────────────────────────────────────
out = "/Users/pratyushkumar/Projects/vendor-mix/Vendor_API_Schema_v1.docx"
doc.save(out)
print(f"Saved: {out}")
